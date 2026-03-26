from docx import Document
import PyPDF2
import os


def extract_text_from_document(file_path):
    """
    Extract text from various document types (PDF, DOCX, TXT)
    
    Args:
        file_path: Path to the document file
        
    Returns:
        Extracted text string
    """
    try:
        if not os.path.exists(file_path):
            return f"File not found: {file_path}"
        
        file_ext = os.path.splitext(file_path)[1].lower()
        
        # Handle DOCX files
        if file_ext == '.docx':
            return extract_from_docx(file_path)
        
        # Handle PDF files
        elif file_ext == '.pdf':
            return extract_from_pdf(file_path)
        
        # Handle TXT files
        elif file_ext == '.txt':
            return extract_from_txt(file_path)
        
        else:
            return f"Unsupported file format: {file_ext}"
    
    except Exception as e:
        return f"Error extracting text: {str(e)}"


def extract_from_docx(file_path):
    """Extract text from DOCX file"""
    try:
        doc = Document(file_path)
        text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
        
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += '\n' + cell.text
        
        return text.strip() if text.strip() else "No text found in document"
    
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"


def extract_from_pdf(file_path):
    """Extract text from PDF file"""
    try:
        text = ""
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            num_pages = len(pdf_reader.pages)
            
            # Extract text from first 20 pages (limit for performance)
            for page_num in range(min(20, num_pages)):
                page = pdf_reader.pages[page_num]
                text += page.extract_text() + '\n'
        
        return text.strip() if text.strip() else "No text found in PDF"
    
    except Exception as e:
        return f"Error reading PDF: {str(e)}"


def extract_from_txt(file_path):
    """Extract text from TXT file"""
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
            text = file.read()
        
        return text.strip() if text.strip() else "No text found in file"
    
    except Exception as e:
        return f"Error reading TXT: {str(e)}"
